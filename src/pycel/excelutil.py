# -*- coding: UTF-8 -*-
#
# Copyright 2011-2019 by Dirk Gorissen, Stephen Rauch and Contributors
# All rights reserved.
# This file is part of the Pycel Library, Licensed under GPLv3 (the 'License')
# You may not use this work except in compliance with the License.
# You may obtain a copy of the Licence at:
#   https://www.gnu.org/licenses/gpl-3.0.en.html

"""
Excel utility functions for address parsing and manipulation.

This module provides utilities for working with Excel addresses,
ranges, and related parsing operations. Evaluation-related code
has been removed for security reasons.
"""

import collections
import itertools as it
import re

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import (
    get_column_letter,
    quote_sheetname,
    range_boundaries as openpyxl_range_boundaries,
)


ERROR_CODES = frozenset(Tokenizer.ERROR_CODES)
DIV0 = '#DIV/0!'
EMPTY = '#EMPTY!'
VALUE_ERROR = '#VALUE!'
NUM_ERROR = '#NUM!'
NA_ERROR = '#N/A'
NAME_ERROR = "#NAME?"
NULL_ERROR = "#NULL!"
REF_ERROR = "#REF!"

R1C1_ROW_RE_STR = r"R(\[-?\d+\]|\d+)?"
R1C1_COL_RE_STR = r"C(\[-?\d+\]|\d+)?"
R1C1_COORD_RE_STR = f"(?P<row>{R1C1_ROW_RE_STR})?(?P<col>{R1C1_COL_RE_STR})?"
R1C1_COORDINATE_RE = re.compile('^' + R1C1_COORD_RE_STR + '$', re.VERBOSE)

R1C1_RANGE_EXPR = f"""
(?P<min_row>{R1C1_ROW_RE_STR})?
(?P<min_col>{R1C1_COL_RE_STR})?
(:(?P<max_row>{R1C1_ROW_RE_STR})?
(?P<max_col>{R1C1_COL_RE_STR})?)?
"""

R1C1_RANGE_RE = re.compile('^' + R1C1_RANGE_EXPR + '$', re.VERBOSE)

TABLE_REF_RE = re.compile(r"^(?P<table_name>[^[]+)\[(?P<table_selector>.*)\]$")

TABLE_SELECTOR_RE = re.compile(
    r"^(?P<row_or_column>[^[]+)$|"
    r"^@\[(?P<this_row_column>[^[]*)\]$|"
    r"^ *(?P<rows>(\[([^\]]+)\] *, *)*)"
    r"(\[(?P<start_col>[^\]]+)\] *: *)?"
    r"(\[(?P<end_col>.+)\] *)?$")

MAX_COL = 16384
MAX_ROW = 1048576

VALID_R1C1_RANGE_ITEM_COMBOS = {
    (0, 1, 0, 1),
    (1, 0, 1, 0),
    (1, 1, 1, 1),
}


AddressSize = collections.namedtuple('AddressSize', 'height width')


class PyCelException(Exception):
    """Base class for PyCel errors"""


class AddressMixin:

    def __str__(self):
        return self.address

    @property
    def has_sheet(self):
        """Does the address have a sheet?"""
        return bool(self.sheet)

    @staticmethod
    def quote_sheet(sheet):
        if ' ' in sheet:
            sheet = quote_sheetname(sheet)
        return sheet

    @property
    def quoted_address(self):
        """requote the sheetname if going to include in formulas"""
        return f"{self.quote_sheet(self.sheet)}!{self.coordinate}"

    @property
    def abs_address(self):
        return f"{self.quote_sheet(self.sheet)}!{self.abs_coordinate}"

    @property
    def sort_key(self):
        return self.sheet, self.col_idx, self.row

    def _union_instersection(self, other, min_, max_):
        """Assumes rectangular only"""
        if not is_address(other):
            other = AddressRange.create(other)
        if self.sheet and other.sheet and self.sheet != other.sheet:
            return VALUE_ERROR

        min_col_idx = min_(self.col_idx, other.col_idx)
        min_row = min_(self.row, other.row)

        max_col_idx = max_(self.col_idx + self.size.width,
                           other.col_idx + other.size.width) - 1
        max_row = max_(self.row + self.size.height,
                       other.row + other.size.height) - 1

        if max_col_idx < min_col_idx or max_row < min_row:
            return NULL_ERROR

        elif max_col_idx == min_col_idx and max_row == min_row:
            return AddressCell((min_col_idx, min_row, max_col_idx, max_row),
                               sheet=self.sheet or other.sheet)
        else:
            return AddressRange((min_col_idx, min_row, max_col_idx, max_row),
                                sheet=self.sheet or other.sheet)

    def __pow__(self, other):
        return self._union_instersection(other, min, max)

    def __rpow__(self, other):
        return self._union_instersection(other, min, max)

    def __and__(self, other):
        return self._union_instersection(other, max, min)

    def __rand__(self, other):
        return self._union_instersection(other, max, min)


class AddressRange(collections.namedtuple(
        'Address', 'address sheet start end coordinate'), AddressMixin):
    """ Helper class for constructing, validating and accessing Range Addresses

    **Tuple Attributes:**

    .. py:attribute:: address

        `AddressRange` as a string

    .. py:attribute:: sheet

        Sheet name

    .. py:attribute:: start

        `AddressCell` for upper left corner of `AddressRange`

    .. py:attribute:: end

        `AddressCell` for lower right corner of `AddressRange`

    .. py:attribute:: coordinate

        Address without the sheetname

    **Non-tuple Attributes:**

    """

    def __new__(cls, address, *args, sheet=''):
        if args:
            return super(AddressRange, cls).__new__(cls, address, *args)

        if isinstance(address, str):
            return cls.create(address, sheet=sheet)

        elif isinstance(address, AddressCell):
            return AddressCell(address, sheet=sheet)

        elif isinstance(address, AddressRange):
            if not sheet or sheet == address.sheet:
                return address

            elif not address.sheet:
                start = AddressCell(address.start.coordinate, sheet=sheet)
                end = AddressCell(address.end.coordinate, sheet=sheet)

            else:
                raise ValueError(f"Mismatched sheets '{address}' and '{sheet}'")

        else:
            assert (isinstance(address, tuple) and 4 == len(address) and
                    None in address or address[0:2] != address[2:]), \
                f"AddressRange expected a range '{address}'"

            start_col, start_row, end_col, end_row = address
            start = AddressCell((start_col, start_row, start_col, start_row), sheet=sheet)
            end = AddressCell((end_col, end_row, end_col, end_row), sheet=sheet)

        coordinate = f'{start.coordinate}:{end.coordinate}'

        format_str = '{0}!{1}' if sheet else '{1}'
        return super(AddressRange, cls).__new__(
            cls, format_str.format(sheet, coordinate),
            sheet, start, end, coordinate)

    def __contains__(self, address):
        address = AddressCell(address)
        return (self.start.row <= address.row <= self.end.row and
                self.start.col_idx <= address.col_idx <= self.end.col_idx)

    @property
    def col_idx(self):
        """col_idx for left column"""
        return self.start.col_idx

    @property
    def row(self):
        """top row"""
        return self.start.row

    @property
    def abs_coordinate(self):
        return f'{self.start.abs_coordinate}:{self.end.abs_coordinate}'

    # Is this address a range?
    is_range = True

    @property
    def is_unbounded_range(self):
        """Is this address an unbounded range?"""
        rows, cols = self.size
        return rows == MAX_ROW or cols == MAX_COL

    @property
    def size(self):
        """Range dimensions"""
        if not hasattr(self, '_size'):
            if 0 in (self.end.row, self.start.row):
                height = MAX_ROW
            else:
                height = self.end.row - self.start.row + 1

            if 0 in (self.end.col_idx, self.start.col_idx):
                width = MAX_COL
            else:
                width = self.end.col_idx - self.start.col_idx + 1

            self._size = AddressSize(height, width)
        return self._size

    @property
    def rows(self):
        """Get each address for every cell, yields one row at a time."""
        col_range = self.start.col_idx, self.end.col_idx + 1
        for row in range(self.start.row, self.end.row + 1):
            yield (AddressCell((col, row, col, row), sheet=self.sheet)
                   for col in range(*col_range))

    @property
    def cols(self):
        """Get each address for every cell, yields one column at a time."""
        col_range = self.start.col_idx, self.end.col_idx + 1
        for col in range(*col_range):
            yield (AddressCell((col, row, col, row), sheet=self.sheet)
                   for row in range(self.start.row, self.end.row + 1))

    def address_at_offset(self, row_inc=0, col_inc=0):
        return self.start.address_at_offset(row_inc=row_inc, col_inc=col_inc)

    @property
    def resolve_range(self):
        """Return nested tuples with an AddressCell for each element"""
        assert not self.is_unbounded_range
        return tuple(tuple(row) for row in self.rows)

    @classmethod
    def create(cls, address, sheet='', cell=None):
        """ Factory method.

        Able to construct R1C1, defined names, and structured references
        style addresses, if passed a `excelcompiler._Cell`.

        :param address: str, AddressRange, AddressCell
        :param sheet: sheet for address, if not included
        :param cell: `excelcompiler._Cell` reference
        :return: `AddressRange or AddressCell`
        """

        if isinstance(address, AddressRange):
            return AddressRange(address, sheet=sheet)

        elif isinstance(address, AddressCell):
            return AddressCell(address, sheet=sheet)

        elif address in ERROR_CODES:
            return address

        sheetname, addr = split_sheetname(address, sheet=sheet)
        addr_tuple, sheetname = range_boundaries(
            addr, sheet=sheetname, cell=cell)

        if isinstance(addr_tuple, AddressMultiAreaRange):
            return addr_tuple
        elif None in addr_tuple or addr_tuple[0:2] != addr_tuple[2:]:
            return AddressRange(addr_tuple, sheet=sheetname)
        else:
            return AddressCell(addr_tuple, sheet=sheetname)


class AddressCell(collections.namedtuple(
        'AddressCell', 'address sheet col_idx row coordinate'), AddressMixin):
    """ Helper class for constructing, validating and accessing Cell Addresses

    **Tuple Attributes:**

    .. py:attribute:: address

        `AddressRange` as a string

    .. py:attribute:: sheet

        Sheet name

    .. py:attribute:: col_idx

        Column number as a 1 based index

    .. py:attribute:: row

        Row number as a 1 based index

    .. py:attribute:: coordinate

        Address without the sheetname

    **Non-tuple Attributes:**

    """

    def __new__(cls, address, *args, sheet=''):
        if args:
            return super(AddressCell, cls).__new__(cls, address, *args)

        if isinstance(address, str):
            return cls.create(address, sheet=sheet)

        elif isinstance(address, AddressCell):
            if not sheet or sheet == address.sheet:
                return address

            elif not address.sheet:
                col_idx, row, coordinate = address[2:5]

            else:
                raise ValueError(f"Mismatched sheets '{address}' and '{sheet}'")

        else:
            assert (isinstance(address, tuple) and 4 == len(address) and
                    None not in address or address[0:2] == address[2:]), \
                f"AddressCell expected a cell '{address}'"

            col_idx, row = (a or 0 for a in address[:2])
            column = (col_idx or '') and get_column_letter(col_idx)
            coordinate = f'{column}{row or ""}'

        if sheet:
            format_str = '{0}!{1}'
        else:
            format_str = '{1}'

        return super(AddressCell, cls).__new__(
            cls, format_str.format(sheet, coordinate),
            sheet, col_idx, row, coordinate)

    def __contains__(self, address):
        return self == AddressCell(address)

    # Is this address a range?
    is_range = False

    # Is this address an unbounded range?"""
    is_unbounded_range = False

    size = AddressSize(1, 1)

    @property
    def column(self):
        """column letter"""
        return (self.col_idx or '') and get_column_letter(self.col_idx)

    def inc_col(self, inc):
        """ Generate an address offset by `inc` columns.

        :param inc: integer number of columns to offset by
        """
        return (self.col_idx + inc - 1) % MAX_COL + 1

    def inc_row(self, inc):
        """ Generate an address offset by `inc` rows.

        :param inc: integer number of rows to offset by
        """
        return (self.row + inc - 1) % MAX_ROW + 1

    @property
    def abs_coordinate(self):
        return f'${self.column}${self.row}'

    def address_at_offset(self, row_inc=0, col_inc=0):
        """ Construct an `AddressCell` offset from the address

        :param row_inc: Number of rows to offset.
        :param col_inc: Number of columns to offset
        :return: `AddressCell`
        """
        new_col = self.inc_col(col_inc)
        new_row = self.inc_row(row_inc)
        return AddressCell((new_col, new_row, new_col, new_row),
                           sheet=self.sheet)

    @property
    def start(self):
        return self

    @property
    def end(self):
        return self

    @property
    def resolve_range(self):
        """Return nested tuples with an AddressCell for each element"""
        return (self, ),

    @classmethod
    def create(cls, address, sheet='', cell=None):
        """ Factory method.

        Able to construct R1C1, defined names, and structured references
        style addresses, if passed a `excelcomppiler._Cell`.

        :param address: str, AddressRange, AddressCell
        :param sheet: sheet for address, if not included
        :param cell: `excelcompiler._Cell` reference
        :return: `AddressCell`
        """
        addr = AddressRange.create(address, sheet=sheet, cell=cell)
        if not isinstance(addr, AddressCell):
            raise ValueError(f"{address} is not a valid coordinate")
        return addr


class AddressMultiAreaRange(tuple):
    """Multi-Area Address Range"""

    def __str__(self):
        return ','.join(str(addr) for addr in self)

    def __contains__(self, address):
        address = AddressCell(address)
        return any(address in addr for addr in self)

    # Is this address a range?
    is_range = True

    @property
    def is_unbounded_range(self):
        """Is this address an unbounded range?"""
        return any(addr.is_unbounded_range for addr in self
                   if isinstance(addr, AddressRange))

    @property
    def resolve_range(self):
        """Return nested tuples with an AddressCell for each element"""
        return it.chain.from_iterable(addr.resolve_range for addr in self)


def is_address(addr):
    return isinstance(addr, (AddressCell, AddressRange))


def unquote_sheetname(sheetname):
    """
    Remove quotes from around, an embedded "''" in, quoted sheetnames

    sheetnames with special characters are quoted in formulas
    This is the inverse of openpyxl.utils.quote_sheetname
    """
    if sheetname.startswith("'") and sheetname.endswith("'"):
        sheetname = sheetname[1:-1].replace("''", "'")
    return sheetname


def split_sheetname(address, sheet=''):
    sh = ''
    if '!' in address:
        sh, address_part = address.split('!', maxsplit=1)

        # Remove redundant sheet references and deal with inner quotes
        redundant_sheet = unquote_sheetname(sh).replace("'", "''")
        address_part = address_part.replace(f"'{redundant_sheet}'!", '')

        if '!' in address_part:
            raise NotImplementedError(f"Non-rectangular formulas: {address}")
        sh = unquote_sheetname(sh)
        address = address_part

        if sh and sheet and sh != sheet:
            raise ValueError(f"Mismatched sheets '{sh}' and '{sheet}'")

    return sheet or sh, address


def structured_reference_boundaries(address, cell=None):
    # Excel reference: https://support.microsoft.com/en-us/office/
    #   Using-structured-references-with-Excel-tables-
    #   F5ED2452-2337-4F71-BED3-C8AE6D2B276E

    match = TABLE_REF_RE.match(address)
    if not match:
        return None

    if cell is None:
        raise PyCelException(f"Must pass cell for Structured Reference {address}")

    name = match.group('table_name')
    table, sheet = cell.excel.table(name)

    if table is None:
        raise PyCelException(f"Table {name} not found for Structured Reference: {address}")

    boundaries = openpyxl_range_boundaries(table.ref)
    assert None not in boundaries

    selector = match.group('table_selector')

    if not selector:
        # all columns and the data rows
        rows, start_col, end_col = None, None, None

    else:
        selector_match = TABLE_SELECTOR_RE.match(selector)
        if selector_match is None:
            raise PyCelException(f"Unknown Structured Reference Selector: {selector}")

        row_or_column = selector_match.group('row_or_column')
        this_row_column = selector_match.group('this_row_column')

        if row_or_column:
            rows = start_col = None
            end_col = row_or_column

        elif this_row_column:
            rows = '#This Row'
            start_col = None
            end_col = this_row_column

        else:
            rows = selector_match.group('rows')
            start_col = selector_match.group('start_col')
            end_col = selector_match.group('end_col')

            if not rows:
                rows = None

            else:
                assert '[' in rows
                rows = [r.split(']')[0] for r in rows.split('[')[1:]]
                if len(rows) != 1:
                    # not currently supporting multiple row selects
                    raise PyCelException(f"Unknown Structured Reference Rows: {address}")

                rows = rows[0]

        if end_col.startswith('#'):
            # end_col collects the single field case
            assert rows is None and start_col is None
            rows = end_col
            end_col = None

        elif end_col.startswith('@'):
            rows = '#This Row'
            end_col = end_col[1:]
            if len(end_col) == 0:
                end_col = start_col

    if rows is None:
        # skip the headers and footers
        min_row = boundaries[1] + (
            table.headerRowCount if table.headerRowCount else 0)
        max_row = boundaries[3] - (
            table.totalsRowCount if table.totalsRowCount else 0)

    else:
        if rows == '#All':
            min_row, max_row = boundaries[1], boundaries[3]

        elif rows == '#Data':
            min_row = boundaries[1] + (
                table.headerRowCount if table.headerRowCount else 0)
            max_row = boundaries[3] - (
                table.totalsRowCount if table.totalsRowCount else 0)

        elif rows == '#Headers':
            min_row = boundaries[1]
            max_row = boundaries[1] + (
                table.headerRowCount if table.headerRowCount else 0) - 1

        elif rows == '#Totals':
            min_row = boundaries[3] - (
                table.totalsRowCount if table.totalsRowCount else 0) + 1
            max_row = boundaries[3]

        elif rows == '#This Row':
            # ::TODO:: If not in a data row, return #VALUE! How to do this?
            min_row = max_row = cell.address.row

        else:
            raise PyCelException(f"Unknown Structured Reference Rows: {rows}")

    if end_col is None:
        # all columns
        min_col_idx, max_col_idx = boundaries[0], boundaries[2]

    else:
        # a specific column
        column_idx = next((idx for idx, c in enumerate(table.tableColumns)
                           if c.name == end_col), None)
        if column_idx is None:
            raise PyCelException(
                f"Column {end_col} not found for Structured Reference: {address}")
        max_col_idx = boundaries[0] + column_idx

        if start_col is None:
            min_col_idx = max_col_idx

        else:
            column_idx = next((idx for idx, c in enumerate(table.tableColumns)
                               if c.name == start_col), None)
            if column_idx is None:
                raise PyCelException(
                    f"Column {start_col} not found for Structured Reference: {address}")
            min_col_idx = boundaries[0] + column_idx

    if min_row > max_row or min_col_idx > max_col_idx:
        raise PyCelException(f"Columns out of order : {address}")

    return (min_col_idx, min_row, max_col_idx, max_row), sheet


def range_boundaries(address, cell=None, sheet=None):
    try:
        # if this is normal reference then just use the openpyxl converter
        boundaries = openpyxl_range_boundaries(address)
        if None not in boundaries or ':' in address:
            return boundaries, sheet
    except ValueError:
        pass

    # test for R1C1 style address
    boundaries = r1c1_boundaries(address, cell=cell, sheet=sheet)
    if boundaries:
        return boundaries

    # Try to see if the is a structured table reference
    boundaries = structured_reference_boundaries(address, cell=cell)
    if boundaries:
        return boundaries

    # Try to see if this is a defined name
    name_addr = cell and cell.excel and cell.excel.defined_names.get(address)
    if name_addr:
        if len(name_addr) == 1:
            return openpyxl_range_boundaries(name_addr[0][0]), name_addr[0][1]
        else:
            return AddressMultiAreaRange(tuple(
                AddressRange(range_alias, sheet=worksheet)
                for range_alias, worksheet in name_addr)), None

    addrs = address.split(':')
    if len(addrs) > 2:
        # Multi colon range resolves to rectangle containing all nodes
        try:
            nodes = tuple(AddressRange.create(addr, cell=cell, sheet=sheet)
                          for addr in addrs)

            min_col_idx = min(n.col_idx for n in nodes)
            max_col_idx = max((n.col_idx + n.size.width - 1) for n in nodes)
            min_row = min(n.row for n in nodes)
            max_row = max((n.row + n.size.height - 1) for n in nodes)

            sheets = {n.sheet for n in nodes if n.sheet}
            if not sheet:
                sheet = next(iter(sheets), None)
            assert not sheets or sheets == {sheet}

            return (min_col_idx, min_row, max_col_idx, max_row), sheet
        except ValueError:
            pass

    raise ValueError(f"{address} is not a valid coordinate or range")


def r1c1_boundaries(address, cell=None, sheet=None):
    """
    R1C1 reference style

    You can also use a reference style where both the rows and the columns on
    the worksheet are numbered. The R1C1 reference style is useful for
    computing row and column positions in macros. In the R1C1 style, Excel
    indicates the location of a cell with an "R" followed by a row number
    and a "C" followed by a column number.

    Reference   Meaning

    R[-2]C      A relative reference to the cell two rows up and in
                the same column

    R[2]C[2]    A relative reference to the cell two rows down and
                two columns to the right

    R2C2        An absolute reference to the cell in the second row and
                in the second column

    R[-1]       A relative reference to the entire row above the active cell

    R           An absolute reference to the current row as part of a range

    """

    # test for R1C1 style address
    m = R1C1_RANGE_RE.match(address)

    if not m:
        return None

    def from_relative_to_absolute(r1_or_c1):
        def require_cell():
            assert cell is not None, \
                f"Must pass a cell to decode a relative address {address}"

        if not r1_or_c1.endswith(']'):
            if len(r1_or_c1) > 1:
                return int(r1_or_c1[1:])

            else:
                require_cell()
                if r1_or_c1[0].upper() == 'R':
                    return cell.row
                else:
                    return cell.col_idx

        else:
            require_cell()
            if r1_or_c1[0].lower() == 'r':
                return (cell.row + int(r1_or_c1[2:-1]) - 1) % MAX_ROW + 1
            else:
                return (cell.col_idx + int(r1_or_c1[2:-1]) - 1) % MAX_COL + 1

    min_col, min_row, max_col, max_row = (
        g if g is None else from_relative_to_absolute(g) for g in (
            m.group(n) for n in ('min_col', 'min_row', 'max_col', 'max_row')
        )
    )

    items_present = (min_col is not None, min_row is not None,
                     max_col is not None, max_row is not None)

    is_range = ':' in address
    if (is_range and items_present not in VALID_R1C1_RANGE_ITEM_COMBOS or
            not is_range and sum(items_present) < 2):
        raise ValueError(f"{address} is not a valid coordinate or range")

    if min_col is not None:
        min_col = min_col

    if min_row is not None:
        min_row = min_row

    if max_col is not None:
        max_col = max_col
    else:
        max_col = min_col

    if max_row is not None:
        max_row = max_row
    else:
        max_row = min_row

    return (min_col, min_row, max_col, max_row), sheet


def flatten(data, coerce=lambda x: x):
    """ flatten items, converting top level items as needed

    :param data: data to flatten
    :param coerce: apply coercion to top level, but not to sub ranges
    :return: flattened (coerced) items
    """
    if isinstance(data, collections.abc.Iterable) and not isinstance(
            data, (str, AddressRange, AddressCell)):
        for item in data:
            yield from flatten(item, coerce=coerce)
    else:
        yield coerce(data)


def uniqueify(seq):
    seen = set()
    return tuple(x for x in seq if x not in seen and not seen.add(x))


def list_like(data):
    return (not isinstance(data, (str, AddressRange, AddressCell)) and
            isinstance(data, collections.abc.Iterable))
